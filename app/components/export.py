"""Branded Excel export — same house style as the Market Data Portal.

Every workbook carries the numbers AND the provenance: which wave, which
filters, which base each percentage was computed on. A spreadsheet of
percentages with no base attached is the thing that ends up in a client deck
saying the wrong number, so the "About this export" sheet is not optional.
"""
from __future__ import annotations

import hashlib
import io
from datetime import datetime
from typing import Callable, Iterable, Optional, Sequence

import pandas as pd
import streamlit as st

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

RED = "D62E2F"
DARK = "2D2A29"
LGRAY = "E0E0E0"
ALT_ROW = "F9F9F9"
HDR_BG = "F0F0F0"


def build_workbook(
    sheets: dict[str, pd.DataFrame],
    title: str,
    notes: Optional[Sequence[tuple[str, str]]] = None,
    percent_columns: Iterable[str] = ("pct", "%", "Share", "share"),
) -> bytes:
    """Build a branded multi-sheet workbook.

    `sheets` maps sheet name -> DataFrame. `notes` becomes the provenance
    sheet: (label, value) pairs such as ("Base", "n=222 — department-store
    buyers").
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color=LGRAY)
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, df in sheets.items():
        ws = wb.create_sheet(str(sheet_name)[:31])

        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=max(len(df.columns), 1))
        head = ws.cell(row=1, column=1, value=title)
        head.font = Font(name="Inter", size=13, bold=True, color="FFFFFF")
        head.fill = PatternFill("solid", start_color=RED, end_color=RED)
        head.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 26

        stamp = ws.cell(row=2, column=1,
                        value=f"Coresight Research · Consumer Insight Portal · "
                              f"{datetime.now():%d %b %Y %H:%M}")
        stamp.font = Font(name="Inter", size=9, italic=True, color="6B6A68")

        header_row = 4
        for col, name in enumerate(df.columns, start=1):
            c = ws.cell(row=header_row, column=col, value=str(name))
            c.font = Font(name="Inter", size=10, bold=True, color=DARK)
            c.fill = PatternFill("solid", start_color=HDR_BG, end_color=HDR_BG)
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[header_row].height = 28

        pct_cols = {
            i for i, name in enumerate(df.columns, start=1)
            if any(tok.lower() in str(name).lower() for tok in percent_columns)
        }

        for r, row in enumerate(df.itertuples(index=False, name=None),
                                start=header_row + 1):
            shade = (r - header_row) % 2 == 0
            for col, value in enumerate(row, start=1):
                if value is None or (isinstance(value, float) and pd.isna(value)):
                    value = None
                c = ws.cell(row=r, column=col, value=value)
                c.font = Font(name="Inter", size=10, color=DARK)
                c.border = border
                if shade:
                    c.fill = PatternFill("solid", start_color=ALT_ROW, end_color=ALT_ROW)
                if col in pct_cols and isinstance(value, (int, float)):
                    c.number_format = "0.0%"
                elif isinstance(value, (int, float)):
                    c.number_format = "#,##0"

        for col, name in enumerate(df.columns, start=1):
            longest = max(
                [len(str(name))] +
                [len(str(v)) for v in df.iloc[:, col - 1].head(200).tolist()]
            )
            ws.column_dimensions[get_column_letter(col)].width = min(max(12, longest + 3), 62)
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # ── provenance ────────────────────────────────────────────────────────
    ws = wb.create_sheet("About this export")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 92
    ws.merge_cells("A1:B1")
    h = ws.cell(row=1, column=1, value="About this export")
    h.font = Font(name="Inter", size=13, bold=True, color="FFFFFF")
    h.fill = PatternFill("solid", start_color=RED, end_color=RED)
    h.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    rows: list[tuple[str, str]] = [
        ("Exported", f"{datetime.now():%d %B %Y, %H:%M}"),
        ("Source", "Coresight Consumer Insight Portal (CSI tables, dwh_stg)"),
    ]
    rows += list(notes or [])
    rows.append((
        "Reading these numbers",
        "Percentages are computed on the base shown for each question, not on the "
        "wave total. Most of this questionnaire is routed, so bases differ by "
        "question — always quote the base alongside the figure.",
    ))
    for i, (label, value) in enumerate(rows, start=3):
        a = ws.cell(row=i, column=1, value=label)
        a.font = Font(name="Inter", size=10, bold=True, color=DARK)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=i, column=2, value=str(value))
        b.font = Font(name="Inter", size=10, color=DARK)
        b.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    buf.close()
    return data


def download_button(
    label: str,
    filename: str,
    builder: Callable[[], bytes],
    key_seed: str = "",
) -> None:
    """Excel download.

    `builder` is passed to Streamlit as a callable so the workbook is only
    assembled when someone actually clicks — building it on every rerun makes
    the page crawl for a file most people never open. `on_click="ignore"`
    keeps the download from triggering a rerun and losing the filter state.
    """
    key = "xl-" + hashlib.md5(f"{filename}|{label}|{key_seed}".encode()).hexdigest()[:12]
    st.download_button(
        label=label,
        data=builder,
        file_name=filename,
        mime=EXCEL_MIME,
        key=key,
        icon=":material/table:",
        on_click="ignore",
        width="content",
    )
