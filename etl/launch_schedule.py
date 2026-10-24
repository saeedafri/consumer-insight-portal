"""The data team's launch plan, checked week by week against what Forsta actually fielded.

The plan is the "Tracking by Week" tab of the data team's workbook (Sector Survey Questions 2022.xlsx): one
row per lead sector, one column per 12-week cycle holding that sector's Monday, and coloured cells under
Inflation / Financial Health / Holiday / GLP-1 / GenAI / Tariffs marking which topics ride with which sector.
Importing it writes config/launch_schedule.yml, so neither the API nor the ETL needs the workbook.

A wave launches Monday about 18:00 IST, fields until its quota of ~400 qualified fills (median 12 h, most
by Tuesday night) and then closes itself; the register (cip_forsta_survey) is refreshed by every ETL run.

    python -m etl.launch_schedule --import "~/Downloads/Sector Survey Questions 2022.xlsx"
    python -m etl.launch_schedule              # last 6 and next 8 weeks, planned vs fielded
    python -m etl.launch_schedule --launchd    # rewrite the launchd plist: a start every Monday 17:30, plus Friday

Past the workbook's last Monday the rule takes over: the tracker launches every Monday, except the dates listed
under `skip:` in config/launch_schedule.yml (kept across imports). So the plan never runs out, and launchd starts
the job every Monday (the run itself decides whether that Monday is a launch: planned_today).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Optional

import yaml
from sqlalchemy import bindparam, text

from app.core.database import get_engine
from etl import survey_calendar

PLAN = Path(__file__).resolve().parents[1] / "config" / "launch_schedule.yml"
LAUNCHD = Path(__file__).resolve().parents[1] / "scripts" / "launchd" / "com.coresight.cip-forsta-etl.plist"
WATCH_FROM = (17, 30)                  # IST on a planned Monday: the Teams heads-up, then the watch starts
CATCH_ALL = {"Weekday": 5, "Hour": 9, "Minute": 0}   # Friday: one short run for annual and ad hoc studies
SHEET = "Tracking by Week"
FIRST_WEEK = date(2025, 7, 1)          # Forsta waves start here; older cycles ran on Qualtrics / SurveyMonkey
RULE_WEEKS = 8                         # how far past today the rule's Mondays are listed
QUOTA = 400                            # qualified completes per weekly wave (400–414 in all 68 Forsta waves)
TOPICS = {"Inflation": "inflation", "Financial Health": "financial health", "Holiday": "holiday",
          "GLP-1": r"glp-?1", "GenAI": "genai", "Tariffs": "tariff"}


def sector_code(name: str) -> Optional[str]:
    name = name.replace("/", " ")
    return next((s["code"] for s in survey_calendar.load_config()["sectors"] if re.search(s["title"], name, re.I)), None)


def read_workbook(path: Path, today: date) -> dict:
    """The plan from the workbook: one entry per Monday. Topics apply from the first Monday after `today` —
    the tab says which topics go with which sector now, not which went with it in past cycles."""
    import openpyxl

    sheet = openpyxl.load_workbook(path, data_only=True)[SHEET]
    header = {c: str(sheet.cell(4, c).value or "").strip() for c in range(1, sheet.max_column + 1)}
    cycles = {c: int(m.group(1)) for c in header if (m := re.match(r"Cycle (\d+)", str(sheet.cell(3, c).value or "")))}
    topic_cols = {c: name for c, name in header.items() if name in TOPICS}
    weeks: dict[date, dict] = {}
    fixed = []
    for r in range(5, sheet.max_row + 1):    # sector rows run down to the first blank row; holidays follow it
        name = " ".join(str(sheet.cell(r, 3).value or "").split())
        if not name:
            break
        topics = [t for c, t in topic_cols.items() if sheet.cell(r, c).fill.fill_type]
        dates = [(cyc, sheet.cell(r, c).value) for c, cyc in sorted(cycles.items())]
        dates = [(cyc, d.date()) for cyc, d in dates if isinstance(d, datetime)]
        for i, (cyc, day) in enumerate(dates):    # a typed year that breaks the row's order (2027-05-11 between
            before = dates[i - 1][1] if i else date.min    # 2026-02-16 and 2026-08-03) is a typo for the year before
            after = dates[i + 1][1] if i + 1 < len(dates) else date.max
            if not before < day < after and before < day.replace(year=day.year - 1) < after:
                fixed.append(f"{name}, cycle {cyc}: {day} → {day.replace(year=day.year - 1)}")
                day = day.replace(year=day.year - 1)
            if day < FIRST_WEEK:
                continue
            week = weeks.setdefault(day, {"date": day, "cycle": cyc, "sectors": [], "names": [], "topics": []})
            week["names"].append(name.split(" (")[0])
            code = sector_code(name)
            if code and code not in week["sectors"]:
                week["sectors"].append(code)
            if day > today:
                week["topics"] += [t for t in topics if t not in week["topics"]]
    return {"source": f"{path.name} → {SHEET}", "imported": today, "corrected": fixed,
            "weeks": [weeks[d] for d in sorted(weeks)]}


def import_workbook(path: Path, today: Optional[date] = None) -> dict:
    plan = {**read_workbook(path, today or date.today()), "skip": load_plan().get("skip") or []}
    head = ("# Generated by `python -m etl.launch_schedule --import <workbook>` — edit the workbook, not this file,\n"
            "# except `skip:` (Mondays with no launch, e.g. a holiday), which an import keeps.\n"
            "# One entry per Monday launch: lead sector(s), the workbook's names, and the topics planned with it.\n"
            "# After the last one, every Monday not in `skip:` is a launch (etl/launch_schedule.plan_weeks).\n")
    PLAN.write_text(head + yaml.safe_dump(plan, sort_keys=False, allow_unicode=True, width=140), encoding="utf-8")
    load_plan.cache_clear()
    return plan


@lru_cache(maxsize=1)
def load_plan() -> dict:
    return yaml.safe_load(PLAN.read_text(encoding="utf-8")) if PLAN.exists() else {"weeks": []}


def plan_weeks(today: Optional[date] = None) -> list[dict]:
    """The workbook's Mondays, then the rule's (every Monday) up to RULE_WEEKS ahead of today; `skip:` dates out."""
    today = today or date.today()
    plan = load_plan()
    skip = {date.fromisoformat(str(d)) for d in plan.get("skip") or []}
    out = [w for w in plan["weeks"] if w["date"] not in skip]
    last = plan["weeks"][-1] if plan["weeks"] else {"date": FIRST_WEEK - timedelta(days=FIRST_WEEK.weekday() + 7), "cycle": 0}
    day = last["date"] + timedelta(days=7 - last["date"].weekday())
    while day <= today + timedelta(weeks=RULE_WEEKS):
        if day not in skip:
            out.append({"date": day, "cycle": last["cycle"], "sectors": [], "names": ["Weekly tracker"], "topics": []})
        day += timedelta(days=7)
    return out


def status(week: dict, survey: Optional[dict], today: date) -> str:
    if survey is None:
        return "planned" if week["date"] >= today else "not_found"
    if survey["load_state"] == "loaded":
        return "loaded"
    if survey["state"] == "live":
        return "in_field"
    if survey["state"] == "closed":
        return "load_failed" if survey["load_state"] == "failed" else "closed"
    return "setting_up"


def weeks(today: Optional[date] = None) -> dict:
    """Every planned Monday with what Forsta did that week: the matching tracker survey (launched within the
    week, else one still being set up for that lead sector), its progress to quota, and anything that
    differs from the plan."""
    today = today or date.today()
    plan = load_plan()
    with get_engine("app").connect() as conn:
        rows = [dict(r) for r in conn.execute(text(
            "SELECT forsta_path, title, state, launched_at, closed_at, qualified_n, load_state, survey_id, quota_json, quota_at"
            " FROM cip_forsta_survey WHERE title LIKE '%% - inc %%'")).mappings()]
        history: dict = {}
        if rows:
            for e in conn.execute(text("SELECT forsta_path, seen_at, state, load_state, qualified_n FROM cip_forsta_event"
                                       " WHERE forsta_path IN :p ORDER BY seen_at, event_id").bindparams(bindparam("p", expanding=True)),
                                  {"p": [r["forsta_path"] for r in rows]}).mappings():
                history.setdefault(e["forsta_path"], []).append({"at": str(e["seen_at"])[:19].replace(" ", "T") + "Z", "state": e["state"],
                                                                 "load_state": e["load_state"], "qualified": e["qualified_n"]})
    for r in rows:
        for k in ("launched_at", "closed_at"):        # SQLite hands back text
            if isinstance(r[k], str):
                r[k] = datetime.fromisoformat(r[k])
        r["lead"] = next((m["code"] for m in survey_calendar.classify(r["title"]) if m["is_lead"]), None)
    used: set[str] = set()
    out = []
    planned = plan_weeks(today)
    for week in planned:
        monday = week["date"]
        free = [r for r in rows if r["forsta_path"] not in used]
        days = lambda r: abs((r["launched_at"].date() - monday).days)   # noqa: E731
        same_sector = [r for r in free if r["launched_at"] and r["lead"] in week["sectors"] and days(r) <= 10]
        same_week = [r for r in free if r["launched_at"] and -3 <= (r["launched_at"].date() - monday).days < 4]
        waiting = [r for r in free if not r["launched_at"] and r["lead"] in week["sectors"] and monday >= today]
        survey = (sorted(same_sector, key=days) or same_week or waiting or [None])[0]   # a wave fielded a few days
                                                                                        # early still counts
        if survey:
            used.add(survey["forsta_path"])
        notes = []
        if survey and week["sectors"] and survey["lead"] not in week["sectors"]:
            notes.append("Different lead sector from the plan")
        missing = [t for t in week["topics"] if survey and not re.search(TOPICS[t], survey["title"], re.I)]
        for topic in missing:                     # the topic's next sector in the rotation picks it up
            later = next((w for w in planned if w["date"] > monday and topic in w["topics"]), None)
            notes.append(f"{topic} not in the survey title" + (
                f"; next planned {later['date']:%b %d} with {' + '.join(later['names'])}" if later else ""))
        out.append({
            "date": str(monday), "cycle": week["cycle"], "sector": " + ".join(week["names"]),
            "topics": week["topics"], "status": status(week, survey, today),
            "survey_id": survey and survey["load_state"] == "loaded" and survey["survey_id"] or None,
            "title": survey and survey["title"], "qualified": survey and survey["qualified_n"], "quota": QUOTA,
            "launched": survey and survey["launched_at"] and survey["launched_at"].isoformat() + "Z",
            "closed": survey and survey["closed_at"] and survey["closed_at"].isoformat() + "Z", "notes": notes,
            "quota_cells": json.loads(survey["quota_json"]) if survey and survey["quota_json"] else None,
            "quota_at": survey and survey["quota_at"] and str(survey["quota_at"])[:19].replace(" ", "T") + "Z",
            "history": history.get(survey["forsta_path"], []) if survey else []})
    last = plan["weeks"][-1]["date"] if plan["weeks"] else None
    topics = []
    for topic, pattern in TOPICS.items():         # each topic's rotation: last fielded, next planned
        fielded = [r for r in rows if r["launched_at"] and re.search(pattern, r["title"], re.I)]
        ahead = [w for w in out if topic in w["topics"] and w["date"] >= str(today)]
        topics.append({"topic": topic, "sectors": sorted({w["sector"] for w in ahead}),
                       "last": str(max(r["launched_at"] for r in fielded).date()) if fielded else None,
                       "fielded": len(fielded), "next": ahead[0]["date"] if ahead else None,
                       "next_sector": ahead[0]["sector"] if ahead else None})
    return {"source": plan.get("source"), "imported": str(plan.get("imported")), "quota": QUOTA,
            "plan_ends": str(last) if last else None, "corrected": plan.get("corrected", []), "topics": topics,
            "weeks": out}


def write_launchd(today: Optional[date] = None) -> list[date]:
    """The Forsta job's start times: every Monday 17:30 IST (on a planned one the run watches until its wave is
    in; on a skipped one it is one short run) plus CATCH_ALL. No dates, so it never runs out and never fires on
    the wrong weekday a year later. Returns the planned Mondays ahead."""
    import plistlib

    today = today or date.today()
    job = plistlib.loads(LAUNCHD.read_bytes())
    job["StartCalendarInterval"] = [{"Weekday": 1, "Hour": WATCH_FROM[0], "Minute": WATCH_FROM[1]}, CATCH_ALL]
    LAUNCHD.write_bytes(plistlib.dumps(job))
    return [w["date"] for w in plan_weeks(today) if w["date"] >= today]


def planned_today(today: Optional[date] = None) -> Optional[dict]:
    today = today or date.today()
    return next((w for w in plan_weeks(today) if w["date"] == today), None)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--import", dest="workbook", type=Path)
    ap.add_argument("--launchd", action="store_true")
    args = ap.parse_args()
    if args.launchd:
        mondays = write_launchd()
        print(f"{LAUNCHD}: every Monday 17:30 + Friday catch-all; next planned {mondays[0] if mondays else 'none'}")
        return 0
    if args.workbook:
        plan = import_workbook(args.workbook.expanduser())
        print(f"{len(plan['weeks'])} launch weeks written to {PLAN}")
        for line in plan["corrected"]:
            print("  corrected year:", line)
        return 0
    today = date.today()
    for w in weeks(today)["weeks"]:
        if -42 <= (date.fromisoformat(w["date"]) - today).days <= 56:
            print(f"{w['date']}  {w['sector'][:34]:34} {','.join(w['topics'])[:30]:30} {w['status']:11} "
                  f"{w['qualified'] or '':>4}  {'; '.join(w['notes'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
