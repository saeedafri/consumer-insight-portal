"""Offline synthetic before/after benchmark for CIP remediation."""
from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import logging
import os
import resource
import socket
import statistics
import sys
import tempfile
import time
from datetime import date, timedelta
from pathlib import Path


def _blocked(*_args, **_kwargs):
    raise RuntimeError("network disabled for CIP benchmark")


socket.socket.connect = _blocked
socket.socket.connect_ex = _blocked
socket.socket.sendto = _blocked
if hasattr(socket.socket, "sendmsg"):
    socket.socket.sendmsg = _blocked
socket.create_connection = _blocked
os.environ["APP_ENV"] = "LOCAL"
os.environ["LOCAL_SQLITE_PATH"] = str(Path(tempfile.gettempdir()) / "cip-benchmark-unused.db")
for _name in (
    "ANTHROPIC_API_KEY", "APP_DB_PASSWORD", "APP_DB_USER", "CIP_NGINX_URL",
    "CIP_OIDC_CLIENT_ID", "CIP_OIDC_CLIENT_SECRET", "CIP_OIDC_ISSUER",
    "CIP_PUBLISH_BACKUP_URL", "EMAIL_PASSWORD", "FORSTA_API_KEY", "IDP_BASE_URL",
    "IDP_TOKEN_URL", "IDP_USERINFO_URL", "OIDC_CLIENT_ID", "OIDC_CLIENT_SECRET",
    "STG_DB_HOST", "STG_DB_NAME", "STG_DB_PASSWORD", "STG_DB_USER", "TEAMS_WEBHOOK_URL",
):
    os.environ[_name] = ""


def percentile(samples: list[float], percentile_value: float) -> float:
    ordered = sorted(samples)
    return ordered[min(len(ordered) - 1, int((len(ordered) - 1) * percentile_value))]


def peak_rss_bytes() -> int:
    value = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return int(value if sys.platform == "darwin" else value * 1024)


def engine_at(path: Path):
    import sqlalchemy as sa
    from app.core import database, schema_upgrade

    engine = sa.create_engine(f"sqlite:///{path}", future=True)

    @sa.event.listens_for(engine, "connect")
    def foreign_keys(connection, _record):
        connection.execute("PRAGMA foreign_keys=ON")

    schema_upgrade.install(engine)
    database._install_sqlite_translation(engine)
    database._engines["app"] = engine
    database._engines["etl"] = engine
    return engine


def digest_rows(engine) -> str:
    from sqlalchemy import text

    digest = hashlib.sha256()
    with engine.connect() as conn:
        for sql in (
            "SELECT survey_id, record_no, forsta_uuid, status_code, status_label, is_qualified "
            "FROM cip_respondent ORDER BY survey_id, record_no",
            "SELECT a.survey_id, r.record_no, f.field_name, a.value_code, a.value_label, a.value_number, a.value_text "
            "FROM cip_answer a JOIN cip_respondent r ON r.respondent_id=a.respondent_id "
            "JOIN cip_field f ON f.field_id=a.field_id ORDER BY a.survey_id, r.record_no, f.field_name",
        ):
            for row in conn.execute(text(sql)):
                digest.update(json.dumps(list(row), default=str, ensure_ascii=False).encode())
    return digest.hexdigest()


def qualtrics(size: int, samples: int) -> dict:
    import openpyxl
    import sqlalchemy as sa
    from etl import qualtrics_export
    from tests.test_phase3 import qualtrics_file

    timings, queries, transactions, hashes, errors = [], [], [], [], 0
    with tempfile.TemporaryDirectory(prefix="cip-final-qualtrics-") as raw:
        temp = Path(raw)
        export = qualtrics_file(temp)
        workbook = openpyxl.load_workbook(export)
        sheet = workbook.active
        templates = [list(row) for row in sheet.iter_rows(min_row=3, values_only=True)]
        sheet.delete_rows(3, sheet.max_row - 2)
        for index in range(size):
            row = list(templates[index % len(templates)])
            row[4] = f"R_synthetic_{index}"
            row[1] = f"2025-05-12 {8 + index % 12:02d}:{index % 60:02d}:00"
            sheet.append(row)
        workbook.save(export)

        for sample in range(samples):
            engine = engine_at(temp / f"qualtrics-{sample}.db")
            counters = {"queries": 0, "transactions": 0}
            sa.event.listen(engine, "before_cursor_execute",
                            lambda *_args, c=counters: c.__setitem__("queries", c["queries"] + 1))
            sa.event.listen(engine, "begin",
                            lambda *_args, c=counters: c.__setitem__("transactions", c["transactions"] + 1))
            started = time.perf_counter()
            try:
                qualtrics_export.ingest_export(export, "2025-05-12")
                hashes.append(digest_rows(engine))
            except Exception:
                errors += 1
                hashes.append("")
            timings.append(time.perf_counter() - started)
            queries.append(counters["queries"])
            transactions.append(counters["transactions"])
            engine.dispose()
    return {
        "mode": "qualtrics",
        "respondents": size,
        "samples": samples,
        "elapsed_samples_s": [round(value, 6) for value in timings],
        "median_s": round(statistics.median(timings), 6),
        "p95_s": round(percentile(timings, 0.95), 6),
        "min_s": round(min(timings), 6),
        "max_s": round(max(timings), 6),
        "query_counts": queries,
        "transaction_counts": transactions,
        "output_sha256": sorted(set(hashes)),
        "peak_rss_bytes": peak_rss_bytes(),
        "errors": errors,
    }


def published_digest(root: Path) -> tuple[str, int]:
    digest, size = hashlib.sha256(), 0
    for path in sorted(root.rglob("*.json.gz")):
        body = gzip.decompress(path.read_bytes())
        digest.update(str(path.relative_to(root)).encode())
        digest.update(body)
        size += len(body)
    return digest.hexdigest(), size


def publisher(waves: int, jobs: int, samples: int) -> dict:
    import sqlalchemy as sa
    from sqlalchemy import text
    from etl import publish, serving
    from tests.test_serving import seed

    with tempfile.TemporaryDirectory(prefix="cip-final-publish-") as raw:
        temp = Path(raw)
        engine = engine_at(temp / "publish.db")
        start_day = date(2026, 1, 5)
        with engine.begin() as conn:
            for index in range(waves):
                sid = 1000 + index
                seed(conn, sid, start_day + timedelta(days=7 * index), qualified=404,
                     xm1_yes=125 + index % 20, xm1_base=404)
        for index in range(waves):
            sid = 1000 + index
            serving.build_wave(sid)
            with engine.begin() as conn:
                options = dict(conn.execute(text(
                    "SELECT option_label, option_id FROM srv_result "
                    "WHERE survey_id=:s AND qcode='XM1' AND dim='total'"
                ), {"s": sid}).all())
                respondents = [row[0] for row in conn.execute(text(
                    "SELECT respondent_id FROM cip_respondent WHERE survey_id=:s ORDER BY record_no"
                ), {"s": sid})]
                conn.execute(text(
                    "INSERT INTO srv_answer "
                    "(survey_id,qcode,item_id,option_id,respondent_id,gender,age_band,generation,income_band,census_region,urbanicity) "
                    "VALUES (:s,'XM1',0,:o,:r,:g,:a,:gen,:i,:c,:u)"
                ), [{
                    "s": sid,
                    "o": options["Yes" if n < 125 + index % 20 else "No"],
                    "r": respondent_id,
                    "g": "Female" if n % 2 else "Male",
                    "a": "18-29" if n % 3 else "30-44",
                    "gen": "GenZ" if n % 3 else "Millennial",
                    "i": "Under $25,000" if n % 2 else "$25,000-$49,999",
                    "c": "South" if n % 2 else "Northeast",
                    "u": "Urban" if n % 2 else "Rural",
                } for n, respondent_id in enumerate(respondents)])

        timings, query_counts, discovery_counts, hashes, sizes, errors = [], [], [], [], [], 0
        for sample in range(samples):
            counters = {"queries": 0, "discovery": 0}

            def count_sql(_conn, _cursor, statement, _parameters, _context, _many):
                normalized = " ".join(statement.lower().split())
                counters["queries"] += 1
                if "select distinct w.survey_family" in normalized and "join srv_wave" in normalized:
                    counters["discovery"] += 1

            sa.event.listen(engine, "before_cursor_execute", count_sql)
            root = temp / f"out-{sample}"
            started = time.perf_counter()
            try:
                built = publish.build(root, jobs=jobs)
                digest, size = published_digest(built)
                hashes.append(digest)
                sizes.append(size)
            except Exception:
                errors += 1
                hashes.append("")
                sizes.append(0)
            timings.append(time.perf_counter() - started)
            query_counts.append(counters["queries"])
            discovery_counts.append(counters["discovery"])
            sa.event.remove(engine, "before_cursor_execute", count_sql)
        engine.dispose()
    return {
        "mode": "publisher",
        "waves": waves,
        "respondents_per_wave": 404,
        "jobs": jobs,
        "samples": samples,
        "elapsed_samples_s": [round(value, 6) for value in timings],
        "median_s": round(statistics.median(timings), 6),
        "p95_s": round(percentile(timings, 0.95), 6),
        "min_s": round(min(timings), 6),
        "max_s": round(max(timings), 6),
        "query_counts": query_counts,
        "discovery_query_counts": discovery_counts,
        "output_sha256": sorted(set(hashes)),
        "result_bytes": sorted(set(sizes)),
        "peak_rss_bytes": peak_rss_bytes(),
        "errors": errors,
    }


def main() -> None:
    logging.basicConfig(level=logging.WARNING, force=True)
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--label", required=True)
    parser.add_argument("--mode", choices=("qualtrics", "publisher"), required=True)
    parser.add_argument("--size", type=int, required=True)
    parser.add_argument("--jobs", type=int, default=1)
    parser.add_argument("--samples", type=int, default=5)
    args = parser.parse_args()
    sys.path.insert(0, str(args.source_root.resolve()))
    result = qualtrics(args.size, args.samples) if args.mode == "qualtrics" else publisher(args.size, args.jobs, args.samples)
    print(json.dumps({"label": args.label, **result}, sort_keys=True))


if __name__ == "__main__":
    main()
