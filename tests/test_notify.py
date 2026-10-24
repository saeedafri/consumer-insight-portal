import smtplib

from etl import notify


class FakeSMTP:
    sent = []

    def __init__(self, host, port, timeout=None):
        self.host, self.port = host, port

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def starttls(self):
        pass

    def login(self, user, password):
        self.user = user

    def send_message(self, message):
        FakeSMTP.sent.append((self.host, self.port, self.user, message))


def test_send_attaches_workbook(monkeypatch):
    monkeypatch.setattr(smtplib, "SMTP", FakeSMTP)
    monkeypatch.setenv("EMAIL_PASSWORD", "x")
    assert notify.send("subject", "body", ["a@coresight.com"], [("w.xlsx", b"PK")], "<p>hi</p>")
    host, port, user, message = FakeSMTP.sent[-1]
    assert (host, port, user) == ("smtp.office365.com", 587, "dataautomation@coresight.com")
    assert message["To"] == "a@coresight.com"
    assert [a.get_filename() for a in message.iter_attachments()] == ["w.xlsx"]
    assert message.get_body(("html",)).get_content().strip() == "<p>hi</p>"


def test_send_without_password_does_nothing(monkeypatch):
    monkeypatch.delenv("EMAIL_PASSWORD", raising=False)
    assert notify.send("subject", "body", ["a@coresight.com"]) is False


def test_recipients_fall_back_to_default():
    assert notify.recipients("tracker") == notify.recipients("default")


def test_failure_emails_and_posts_to_teams_once(monkeypatch, tmp_path):
    sent, posted = [], []
    monkeypatch.setattr(notify, "STATE", tmp_path / "state.json")
    monkeypatch.setattr(notify, "send", lambda subject, body, to, attachments=(), page="": sent.append(subject) or True)
    monkeypatch.setattr(notify, "teams", posted.append)
    assert notify.failure("survey X did not load", "boom")
    assert notify.failure("survey X did not load", "boom") is False          # same failure inside the window
    assert sent == ["CIP pipeline FAILED: survey X did not load"]
    assert posted == ["\u274c CIP pipeline failed: survey X did not load. boom (full detail in the email and logs/forsta-etl.log)"]


def test_a_failure_posted_only_to_teams_still_counts_as_told(monkeypatch, tmp_path):
    monkeypatch.setattr(notify, "STATE", tmp_path / "state.json")
    monkeypatch.setattr(notify, "send", lambda *a, **k: False)            # no mail password on this machine
    monkeypatch.setattr(notify, "teams", lambda headline: True)
    assert notify.failure("run crashed", "trace") and notify.already_told("CIP pipeline FAILED: run crashed")


def test_raw_data_is_forstas_labelled_export_with_its_datamap():
    import io
    import openpyxl
    from etl import crosstab_export

    class Forsta:
        def survey_datamap(self, path):
            return {"variables": [
                {"label": "uuid", "type": "text"},
                {"label": "q1r1", "type": "multiple", "rowTitle": "Gone to a bar", "colTitle": None},
                {"label": "D1", "type": "single", "values": [{"value": 1, "title": "Male"}, {"value": 2, "title": "Female"}]},
                {"label": "D2", "type": "single", "values": [{"value": 20, "title": "36"}]},
                {"label": "D9", "type": "single", "values": []}, {"label": "qtime", "type": "float"}]}

        def survey_data(self, path, cond):
            assert cond == "qualified"
            return [{"uuid": "a1", "q1r1": "1", "D1": "2", "D2": "20", "qtime": "143.945"}, {"uuid": "b2", "q1r1": "0", "D1": "1"}]

        def _get(self, endpoint, format):
            assert endpoint.endswith("/datamap") and format == "text"
            return type("R", (), {"content": "[D1]: What is your gender?\nValues: 1-2\n\t1\tKohl’s\n".encode()})()

    book = openpyxl.load_workbook(io.BytesIO(crosstab_export.raw_data(Forsta(), "selfserve/58f/1")))
    assert [list(r) for r in book["A1"].iter_rows(values_only=True)] == [
        ["uuid", "q1r1", "D1", "D2", "D9", "qtime"], ["a1", "Gone to a bar", "Female", 36, None, 143.945],
        ["b2", "NO TO: Gone to a bar", "Male", None, None, None]]
    assert [list(r) for r in book["Datamap"].iter_rows(values_only=True)] == [
        ["[D1]: What is your gender?", None, None], ["Values: 1-2", None, None], [None, 1, "Kohl’s"]]
