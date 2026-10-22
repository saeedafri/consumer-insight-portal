"""The launch plan: read from the data team's workbook, then matched week by week to the Forsta register."""
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from sqlalchemy import text

from etl import launch_schedule


@pytest.fixture()
def workbook(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = launch_schedule.SHEET
    for col, cycle in ((4, 1), (5, 2), (6, 3)):
        sheet.cell(3, col, f"Cycle {cycle}")
    for col, name in enumerate(["", "Week", "Sector", "Date", "Date", "Date", "Inflation", "GenAI"], 1):
        sheet.cell(4, col, name)
    rows = [("Off-price", [datetime(2026, 8, 3), datetime(2027, 10, 26), datetime(2027, 1, 18)], [7]),
            ("Dollar stores", [datetime(2026, 8, 3), datetime(2026, 10, 26), datetime(2027, 1, 18)], []),
            ("Apparel", [datetime(2026, 8, 10), datetime(2026, 11, 2), "N/A"], [8])]
    for r, (name, dates, topics) in enumerate(rows, 5):
        sheet.cell(r, 3, name)
        for col, value in enumerate(dates, 4):
            sheet.cell(r, col, value)
        for col in topics:
            sheet.cell(r, col).fill = PatternFill("solid", fgColor="FF00B050")
    sheet.cell(9, 3, "Christmas/Holiday")          # after the blank row: holidays, not sectors
    sheet.cell(9, 4, datetime(2026, 12, 25))
    path = tmp_path / "plan.xlsx"
    book.save(path)
    return path


def test_workbook_merges_weeks_fixes_typed_years_and_plans_topics_ahead(workbook):
    plan = launch_schedule.read_workbook(workbook, today=date(2026, 10, 21))
    by_day = {w["date"]: w for w in plan["weeks"]}
    assert sorted(by_day) == [date(2026, 8, 3), date(2026, 8, 10), date(2026, 10, 26), date(2026, 11, 2), date(2027, 1, 18)]
    assert by_day[date(2026, 10, 26)]["names"] == ["Off-price", "Dollar stores"]
    assert by_day[date(2026, 10, 26)]["sectors"] == ["OFF_PRICE"]
    assert plan["corrected"] == ["Off-price, cycle 2: 2027-10-26 → 2026-10-26"]
    assert by_day[date(2026, 10, 26)]["topics"] == ["Inflation"]
    assert by_day[date(2026, 8, 3)]["topics"] == []            # topics are the plan from now on, not history
    assert by_day[date(2026, 11, 2)]["sectors"] == ["APPAREL"] and by_day[date(2026, 11, 2)]["topics"] == ["GenAI"]


def test_weeks_match_forsta_and_flag_a_missing_topic(workbook, csi_db, monkeypatch, tmp_path):
    monkeypatch.setattr(launch_schedule, "PLAN", tmp_path / "launch_schedule.yml")
    launch_schedule.import_workbook(workbook, today=date(2026, 10, 21))
    with csi_db.begin() as conn:
        conn.execute(text(
            "INSERT INTO cip_forsta_survey (forsta_path, title, state, launched_at, closed_at, qualified_n, load_state,"
            " survey_id, last_seen) VALUES (:p, :t, :s, :l, :c, :q, :ls, :sid, CURRENT_TIMESTAMP)"), [
            {"p": "a", "t": "Shopping and Spending - inc Off Price + Dollar Stores", "s": "closed",
             "l": datetime(2026, 8, 3, 12, 30), "c": datetime(2026, 8, 4, 1), "q": 402, "ls": "loaded", "sid": 7},
            {"p": "b", "t": "Shopping and Spending - inc Off-price / Dollar Stores + Inflation", "s": "live",
             "l": datetime(2026, 10, 26, 12, 30), "c": None, "q": 180, "ls": "testing", "sid": None},
            {"p": "c", "t": "Shopping and Spending - inc Apparel", "s": "testing",
             "l": None, "c": None, "q": 0, "ls": "testing", "sid": None}])
    weeks = {w["date"]: w for w in launch_schedule.weeks(today=date(2026, 10, 27))["weeks"]}
    assert weeks["2026-08-03"]["status"] == "loaded" and weeks["2026-08-03"]["survey_id"] == 7
    assert weeks["2026-08-10"]["status"] == "not_found"
    assert weeks["2026-10-26"]["status"] == "in_field" and weeks["2026-10-26"]["notes"] == []
    assert weeks["2026-11-02"]["status"] == "setting_up"
    assert weeks["2026-11-02"]["notes"] == ["GenAI not in the survey title"]
    assert weeks["2027-01-18"]["status"] == "planned"


def test_fieldwork_ends_when_no_quota_is_left_or_the_time_limit_without_touching_forsta():
    from datetime import timedelta
    from etl.forsta_etl import FIELD_HOURS, fielded, quota_cells

    class Quota:   # Forsta's quota shape (19 Oct 2026 wave): Regions and Gender count; the income flow is conditional
        def __init__(self, south, male):
            self.south, self.male = south, male

        def survey_quota(self, path):
            return {"defines": {"S": {"description": "South"}, "W": {"description": "West"}, "m": {"description": "Male"},
                                "f": {"description": "Female"}},
                    "markers": {"/R/S": {"limit": 1200, "complete": self.south}, "/R/W": {"limit": 800, "complete": 800},
                                "/G/m": {"limit": 1000, "complete": self.male}, "/G/f": {"limit": 1000, "complete": 1000},
                                "/F/x": {"limit": 50, "complete": 3}},
                    "sheets": {"Regions": [{"sheetCond": None, "description": "Regions", "cells": [{"marker": "/R/S", "components": ["S"]}, {"marker": "/R/W", "components": ["W"]}]}],
                               "Gender": [{"sheetCond": None, "description": "Gender", "cells": [{"marker": "/G/m", "components": ["m"]}, {"marker": "/G/f", "components": ["f"]}]}],
                               "FLOWINC": [{"sheetCond": "condition.F", "description": "FLOWINC", "cells": [{"marker": "/F/x", "components": ["x"]}]}]}}

    open_cells = quota_cells(Quota(south=1100, male=990), "s/2")
    assert open_cells[0] == {"sheet": "Regions", "conditional": False, "cell": "South", "limit": 1200, "complete": 1100}
    full = quota_cells(Quota(south=1200, male=1000), "s/2")
    launch = datetime(2026, 10, 26, 12, 30)
    tracker = {"path": "s/1", "title": "Shopping and Spending - inc Apparel", "dateLaunched": "2026-10-26T12:30:00Z"}
    adhoc = {"path": "s/2", "title": "Online Grocery 2026", "dateLaunched": "2026-10-26T12:30:00Z"}
    soon = launch + timedelta(hours=2)
    assert not fielded({**tracker, "qualified": 399}, None, soon)
    assert fielded({**tracker, "qualified": 400}, None, soon)               # Forsta stops trackers at 400, cells short
    assert fielded({**tracker, "qualified": 120}, None, launch + timedelta(hours=FIELD_HOURS))
    assert not fielded({**adhoc, "qualified": 1990}, open_cells, soon)      # quota left in South and Male
    assert fielded({**adhoc, "qualified": 2000}, full, soon)                # no quota left: the conditional flow is ignored
    assert fielded({**adhoc, "qualified": 2000}, open_cells, soon)          # a whole sheet's target reached
