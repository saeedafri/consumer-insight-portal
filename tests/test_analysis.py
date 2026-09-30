"""Cohort engine and export, against a loaded local database.

Skipped unless data/csi_local.db exists — build it with
`bash scripts/run_local.sh` or `python scripts/init_db.py` + the pipeline.
"""
from __future__ import annotations

import os
from pathlib import Path

import pytest

DB = Path(__file__).resolve().parents[1] / "data" / "csi_local.db"
pytestmark = pytest.mark.skipif(not DB.exists(), reason="no local database built")

os.environ.setdefault("APP_ENV", "LOCAL")
os.environ.setdefault("LOCAL_SQLITE_PATH", "data/csi_local.db")


@pytest.fixture(scope="module")
def repo():
    import app.data.repository as _repo
    # strip Streamlit's cache decorator so the functions are callable off-runtime
    for name in dir(_repo):
        fn = getattr(_repo, name)
        if hasattr(fn, "__wrapped__"):
            setattr(_repo, name, fn.__wrapped__)
    return _repo


@pytest.fixture(scope="module")
def qid(repo):
    lookup = repo.question_lookup(1).set_index("qcode")
    return lambda code: int(lookup.loc[code, "question_id"])


# ── the bug this guards ────────────────────────────────────────────────────
def test_single_punch_answers_keep_distinct_codes(repo, qid):
    """The export is in LABEL format, so a single-punch cell holds "Yes", not 1.

    Run through the multi-punch rule, every answered cell becomes code 1 and
    "BNPL users" silently becomes all 404 respondents instead of 133. The
    loader resolves the label back through the datamap; this is the guard.
    """
    dist = repo.analyse(1, qid("BN1"), ())
    counts = dict(zip(dist["answer"], dist["n"]))
    assert counts["Yes"] == 133
    assert counts["No"] == 248
    assert len(counts) == 4, "all four answer options must survive the load"


# ── cohorts reproduce the bases Forsta published ───────────────────────────
@pytest.mark.parametrize("code,expected", [("BN1", 133), ("DP1", 222), ("GP1", 71)])
def test_gate_question_cohorts_match_published_bases(repo, qid, code, expected):
    cohort = ({"kind": "code", "question_id": qid(code), "ids": [1]},)
    assert repo.cohort_size(1, cohort) == expected


def test_demographic_cohorts_match_the_crosstab_summary(repo):
    genz = ({"kind": "profile", "dimension": "generation", "values": ["GenZ"]},)
    assert repo.cohort_size(1, genz) == 54          # Summary sheet: GenZ n=54
    both = ({"kind": "profile", "dimension": "generation",
             "values": ["GenZ", "Millennial"]},)
    assert repo.cohort_size(1, both) == 54 + 159


def test_criteria_intersect_rather_than_union(repo, qid):
    bnpl = {"kind": "code", "question_id": qid("BN1"), "ids": [1]}
    young = {"kind": "profile", "dimension": "generation", "values": ["GenZ"]}
    a, b = repo.cohort_size(1, (bnpl,)), repo.cohort_size(1, (young,))
    both = repo.cohort_size(1, (bnpl, young))
    assert both <= min(a, b)


# ── unfiltered results must still equal the published table ────────────────
def test_unfiltered_analysis_matches_the_published_crosstab(repo, qid):
    df = repo.analyse(1, qid("q4"), ())
    walmart = df[df.answer == "Walmart"].iloc[0]
    assert int(walmart.n) == 241
    assert int(walmart.base_n) == 404
    assert abs(float(walmart.pct) - 0.59653465) < 1e-8


def test_routed_question_keeps_its_own_base_after_filtering(repo, qid):
    """DP2 is asked of department-store buyers only. Filtering on generation
    must not silently widen the denominator back to the cohort."""
    cohort = ({"kind": "profile", "dimension": "generation", "values": ["Millennial"]},)
    n = repo.cohort_size(1, cohort)
    df = repo.analyse(1, qid("DP2"), cohort)
    assert not df.empty
    assert int(df.base_n.max()) < n


def test_break_by_splits_without_losing_the_total(repo, qid):
    total = repo.analyse(1, qid("q4"), ())
    split = repo.analyse(1, qid("q4"), (), "generation")
    t = int(total[total.answer == "Walmart"].n.iloc[0])
    s = int(split[split.answer == "Walmart"].n.sum())
    assert t == s


# ── export ─────────────────────────────────────────────────────────────────
def test_excel_export_is_a_readable_workbook_with_provenance(repo, qid):
    import io
    import openpyxl
    from app.components import export

    df = repo.analyse(1, qid("BN1"), ())
    table = df.rename(columns={"answer": "Answer", "n": "Respondents",
                               "base_n": "Base", "pct": "%"})
    data = export.build_workbook({"Analysis": table}, "BN1 — BNPL usage",
                                 [("Wave", "CSI-US · 2026-09"), ("Base", "n=404")])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Analysis", "About this export"]
    ws = wb["Analysis"]
    assert ws.cell(row=1, column=1).value.startswith("BN1")
    assert ws.cell(row=4, column=2).value == "Answer"
    # percentages must be formatted as percentages, not raw decimals
    pct_col = [c + 1 for c, name in enumerate(table.columns) if name == "%"][0]
    assert ws.cell(row=5, column=pct_col).number_format == "0.0%"
    about = wb["About this export"]
    labels = [about.cell(row=r, column=1).value for r in range(3, 12)]
    assert "Wave" in labels and "Base" in labels


# ── 09/28 wave: routed multi-select items and non-breaking-space labels ─────
def _wave2_qid(repo, code):
    lookup = repo.question_lookup(2).set_index("qcode")
    if code not in lookup.index:
        pytest.skip("09/28 wave not loaded as survey 2")
    return int(lookup.loc[code, "question_id"])


def test_multi_select_base_is_everyone_answering_the_question(repo):
    """BT8 shows cosmetics only to cosmetics buyers; Forsta still divides by
    everyone who answered BT8 (264), not by those shown the item (83)."""
    df = repo.analyse(2, _wave2_qid(repo, "BT8"), ())
    mascara = df[df.answer == "Mascara"].iloc[0]
    assert int(mascara.base_n) == 264
    assert abs(float(mascara.pct) - 0.1515) < 0.0005


def test_every_single_punch_answer_resolves_to_a_code(repo):
    df = repo.analyse(2, _wave2_qid(repo, "HX1"), ())
    assert len(df) == 7 and int(df.n.sum()) == 403


def test_grid_rows_are_separate_distributions(repo):
    df = repo.analyse(2, _wave2_qid(repo, "BT14"), ())
    per_row = df.assign(item=df.answer.str.split(" — ").str[0]).groupby("item").pct.sum()
    assert len(per_row) > 1
    assert ((per_row - 1).abs() < 1e-9).all(), "each rated row must sum to 100%"
